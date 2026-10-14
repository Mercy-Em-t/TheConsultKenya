import os
import win32com.client
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_srs_docx(md_path, docx_path):
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles
    title_style = doc.styles['Title']
    title_style.font.name = 'Calibri'
    title_style.font.size = Pt(26)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
    
    h1_style = doc.styles['Heading 1']
    h1_style.font.name = 'Calibri'
    h1_style.font.size = Pt(18)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(16, 44, 87)
    
    h2_style = doc.styles['Heading 2']
    h2_style.font.name = 'Calibri'
    h2_style.font.size = Pt(14)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(53, 95, 142)
    
    h3_style = doc.styles['Heading 3']
    h3_style.font.name = 'Calibri'
    h3_style.font.size = Pt(12)
    h3_style.font.bold = True
    h3_style.font.color.rgb = RGBColor(80, 80, 80)
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_lines = []
    in_code = False
    code_lines = []
    
    for line in lines:
        raw_line = line.rstrip('\r\n')
        
        # Code blocks
        if raw_line.startswith('```'):
            if in_code:
                in_code = False
                p = doc.add_paragraph('\n'.join(code_lines))
                p.style = 'Normal'
                p.paragraph_format.left_indent = Inches(0.4)
                for r in p.runs:
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(40, 40, 40)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            continue
            
        if in_code:
            code_lines.append(raw_line)
            continue
            
        # Tables
        if '|' in raw_line and (raw_line.strip().startswith('|') or raw_line.strip().endswith('|')):
            in_table = True
            table_lines.append(raw_line)
            continue
        elif in_table:
            # Process table
            in_table = False
            parsed_rows = []
            for tl in table_lines:
                # Skip divider line like |:---|:---|
                if '---' in tl:
                    continue
                cells = [c.strip() for c in tl.strip().split('|')[1:-1]]
                if cells:
                    parsed_rows.append(cells)
            
            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_data in enumerate(parsed_rows):
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < num_cols:
                            cell = tbl.cell(r_idx, c_idx)
                            cell.text = cell_value
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.name = 'Calibri'
                                    r.font.size = Pt(9.5)
                                    if r_idx == 0:
                                        r.font.bold = True
                                        r.font.color.rgb = RGBColor(255, 255, 255)
                            if r_idx == 0:
                                set_cell_background(cell, "102C57") # Navy header
                            elif r_idx % 2 == 1:
                                set_cell_background(cell, "F8F9FA")
                doc.add_paragraph() # space after table
            table_lines = []
            
        # Headings & Text
        stripped = raw_line.strip()
        if not stripped:
            continue
            
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(stripped[2:], style='List Bullet')
        elif stripped.startswith('1. ') or stripped.startswith('2. ') or stripped.startswith('3. '):
            p = doc.add_paragraph(stripped[3:], style='List Number')
        else:
            p = doc.add_paragraph(stripped)
            
    doc.save(docx_path)
    print(f"Saved DOCX to {docx_path}")

def convert_docx_to_pdf(docx_path, pdf_path):
    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)
    
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(abs_docx)
        # 17 = wdFormatPDF
        doc.SaveAs(abs_pdf, FileFormat=17)
        doc.Close()
        print(f"Successfully compiled PDF to {abs_pdf}")
    finally:
        word.Quit()

if __name__ == "__main__":
    md_file = "CONSULT_KENYA_IEEE_SRS_V2.md"
    docx_file = "CONSULT_KENYA_IEEE_SRS_V2.docx"
    pdf_file = "CONSULT_KENYA_IEEE_SRS_V2.pdf"
    
    create_srs_docx(md_file, docx_file)
    convert_docx_to_pdf(docx_file, pdf_file)
